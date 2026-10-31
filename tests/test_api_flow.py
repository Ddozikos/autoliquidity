import os
import tempfile
import time
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DB_PATH = tempfile.mktemp(prefix='autoliquidity_test_', suffix='.db', dir=ROOT / 'data')
os.environ['DATABASE_URL'] = f'sqlite:///{DB_PATH}'
os.environ['APIAD_ENABLED'] = 'false'
for _name in ('AVITO_PRICE_API_URL','AUTORU_PRICE_API_URL','DROM_PRICE_API_URL','SBERAUTO_PRICE_API_URL','VIN_API_URL'):
    os.environ[_name] = ''

from app.main import app  # noqa: E402


class ApiFlowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fixtures=(ROOT/'samples'/'Пример ЗРД.html',ROOT/'samples'/'fssp_100_dokumentov'/'html'/'099_Toyota_Avensis.html')
        if not all(path.exists() for path in fixtures):
            raise unittest.SkipTest('user-provided reference documents are not included in the distributable archive')
        cls.client_context = TestClient(app)
        cls.client = cls.client_context.__enter__()
        login = cls.client.post('/api/login', json={'email':'mvp-test@example.invalid','password':'test'})
        if login.status_code != 200:
            raise AssertionError(login.text)
        cls.headers = {'Authorization':'Bearer '+login.json()['access_token']}

    @classmethod
    def tearDownClass(cls):
        cls.client_context.__exit__(None, None, None)
        from app.database import engine
        engine.dispose()
        for suffix in ('','-wal','-shm'):
            try:
                Path(DB_PATH+suffix).unlink()
            except FileNotFoundError:
                pass

    def _upload_and_wait(self, filename, content, expected_count, expected_type):
        response = self.client.post(
            '/api/process', headers=self.headers,
            data={'min_price':'300000','min_demand':'70'},
            files=[('files',(filename,content,'text/html'))],
        )
        self.assertEqual(response.status_code,200,response.text)
        report_id=response.json()['report_id']
        for _ in range(200):
            report=self.client.get(f'/api/reports/{report_id}',headers=self.headers).json()
            if report['status']!='processing':
                break
            time.sleep(.05)
        self.assertEqual(report['status'],'ready',report)
        self.assertEqual(len(report['cars']),expected_count)
        self.assertEqual({car['doc_type'] for car in report['cars']},{expected_type})
        return report

    def test_upload_history_isolation_and_selected_xlsx(self):
        reference=(ROOT/'samples'/'Пример ЗРД.html').read_bytes()
        report=self._upload_and_wait('Пример ЗРД.html',reference,12,'ban')
        self.assertEqual(report['cars'][0]['vin'],'5TDDKRFH30S008927')
        self.assertIsNone(report['cars'][0]['price'])
        self.assertEqual(report['cars'][0]['market_status'],'not_configured')
        self.assertEqual(report['prioritized_ids'],[])
        self.assertEqual(report['cars'][0]['price'],None)
        self.assertEqual(report['cars'][0]['demand'],None)
        self.assertIsNone(report['cars'][0]['market_checked_at'])

        market=self.client.post('/api/market/search',headers=self.headers,json={'make':'Audi','model':'Q5','year':2019})
        self.assertEqual(market.status_code,200,market.text)
        self.assertEqual(market.json()['status'],'not_configured')
        self.assertIsNone(market.json()['average_price'])
        self.assertEqual(market.json()['offers'],[])

        removal_path=ROOT/'samples'/'fssp_100_dokumentov'/'html'/'099_Toyota_Avensis.html'
        removal=self._upload_and_wait(removal_path.name,removal_path.read_bytes(),1,'removal')
        self.assertEqual(removal['cars'][0]['vin'],'JTNBE46KX73123456')
        from app.database import SessionLocal
        from app.models import Car
        db=SessionLocal()
        try:
            removed_car=db.get(Car,removal['cars'][0]['id']);removed_car.price=1_200_000;removed_car.demand=95
            db.commit()
        finally:db.close()
        removal_detail=self.client.get(f"/api/reports/{removal['id']}",headers=self.headers).json()
        self.assertEqual(removal_detail['prioritized_ids'],[])
        removal_export=self.client.get(f"/api/reports/{removal['id']}/export",headers=self.headers,params={'ids':str(removal['cars'][0]['id'])})
        removal_book=load_workbook(BytesIO(removal_export.content),read_only=True)
        self.assertEqual(removal_book.active.max_row,1)

        renamed=self.client.patch(f"/api/reports/{report['id']}",headers=self.headers,json={'name':'Эталонный пакет'})
        self.assertEqual(renamed.status_code,200)
        history=self.client.get('/api/reports',headers=self.headers).json()
        renamed_item=next(item for item in history if item['name']=='Эталонный пакет')
        self.assertTrue(renamed_item['created_at'].endswith('+00:00'))
        self.assertTrue(renamed_item['created_at'])

        car_id=report['cars'][0]['id']
        xlsx=self.client.get(f"/api/reports/{report['id']}/export",headers=self.headers,params={'ids':str(car_id)})
        self.assertEqual(xlsx.status_code,200)
        workbook=load_workbook(BytesIO(xlsx.content),read_only=True)
        columns=[cell.value for cell in next(workbook.active.iter_rows(min_row=1,max_row=1))]
        self.assertIn('Confidence',columns)
        self.assertIn('Тип документа',columns)

        db=SessionLocal()
        try:
            first=db.get(Car,report['cars'][0]['id']);first.price=1_000_000;first.demand=80;first.motivation_index=2
            second=db.get(Car,report['cars'][1]['id']);second.price=1_500_000;second.demand=75;second.motivation_index=3
            at_minimum=db.get(Car,report['cars'][2]['id']);at_minimum.price=1_000_000;at_minimum.demand=70;at_minimum.motivation_index=1
            low_demand=db.get(Car,report['cars'][3]['id']);low_demand.price=2_000_000;low_demand.demand=69;low_demand.motivation_index=4
            db.commit()
        finally:db.close()
        prioritized=self.client.get(f"/api/reports/{report['id']}",headers=self.headers).json()['prioritized_ids']
        self.assertEqual(prioritized,[second.id,first.id,at_minimum.id])
        selected_ids=','.join(str(car['id']) for car in report['cars'])
        qualified_export=self.client.get(f"/api/reports/{report['id']}/export",headers=self.headers,params={'ids':selected_ids})
        self.assertEqual(qualified_export.status_code,200)
        qualified_book=load_workbook(BytesIO(qualified_export.content),read_only=True)
        self.assertEqual(qualified_book.active.max_row,4)  # header + 3 cars at demand >= 70

        another=self.client.post('/api/login',json={'email':'other-mvp-test@example.invalid','password':'test'}).json()['access_token']
        forbidden=self.client.get(f"/api/reports/{report['id']}",headers={'Authorization':'Bearer '+another})
        self.assertEqual(forbidden.status_code,404)

    def test_two_50_file_batches_are_processed_and_temp_uploads_removed(self):
        import app.main as main_module
        sample=(ROOT/'samples'/'fssp_100_dokumentov'/'html'/'099_Toyota_Avensis.html').read_bytes()
        created=[]
        original=main_module.tempfile.NamedTemporaryFile
        def capture_tempfile(*args,**kwargs):
            handle=original(*args,**kwargs);created.append(Path(handle.name));return handle
        report_ids=[]
        with patch('app.main.tempfile.NamedTemporaryFile',side_effect=capture_tempfile):
            for batch in range(2):
                files=[('files',(f'batch{batch+1}-{index}.html',sample,'text/html')) for index in range(50)]
                response=self.client.post('/api/process',headers=self.headers,data={'min_price':'300000','min_demand':'70'},files=files)
                self.assertEqual(response.status_code,200,response.text)
                report_ids.append(response.json()['report_id'])
            reports=[]
            for _ in range(400):
                reports=[self.client.get(f'/api/reports/{rid}',headers=self.headers).json() for rid in report_ids]
                if all(report['status']!='processing' for report in reports):break
                time.sleep(.05)
        self.assertEqual([report['status'] for report in reports],['ready','ready'])
        self.assertEqual([len(report['cars']) for report in reports],[50,50])
        self.assertTrue(created)
        self.assertTrue(all(not path.exists() for path in created))


if __name__=='__main__':
    unittest.main()
